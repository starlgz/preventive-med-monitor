import pytest
import os
import tempfile
import docx
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors

from app.parsers.column_mapper import ColumnMapper
from app.parsers.word_parser import WordJobParser
from app.parsers.pdf_parser import PdfJobParser
from app.parsers.html_table_parser import HtmlTableJobParser
from app.parsers.dispatcher import ParserDispatcher

def test_column_mapper_enhanced_synonyms():
    """测试新增的用人部门、招录计划、所需专业与代码等同义词映射"""
    headers = ["用人部门", "招录计划", "所需专业与代码", "最低学历学位", "咨询电话"]
    mapping = ColumnMapper.map_columns(headers)
    assert mapping["unit_name"] == 0
    assert mapping["headcount"] == 1
    assert mapping["major"] == 2
    assert mapping["education"] == 3

def test_docx_parser_and_total_filtering():
    """测试 Word .docx 表格解析与合计/统计行自动剔除、单位向上填充"""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tf:
        docx_path = tf.name

    doc = docx.Document()
    table = doc.add_table(rows=4, cols=5)
    
    # 填充表头
    headers = ["用人单位", "招聘岗位", "招考人数", "专业要求", "学历"]
    for i, h in enumerate(headers):
        table.cell(0, i).text = h
        
    # 行 1: 正常岗位
    row1 = ["杭州市疾病预防控制中心", "公卫流调岗", "2", "预防医学（100401K）", "本科及以上"]
    for i, val in enumerate(row1):
        table.cell(1, i).text = val
        
    # 行 2: 空单位（应向上继承）
    row2 = ["", "理化检验岗", "1", "卫生检验与检疫", "硕士研究生"]
    for i, val in enumerate(row2):
        table.cell(2, i).text = val
        
    # 行 3: 合计行（应被过滤）
    row3 = ["合计", "合计", "3", "", ""]
    for i, val in enumerate(row3):
        table.cell(3, i).text = val
        
    doc.save(docx_path)
    
    try:
        jobs = WordJobParser.parse_file(docx_path)
        assert len(jobs) == 2
        assert jobs[0]["unit_name"] == "杭州市疾病预防控制中心"
        assert jobs[0]["job_name"] == "公卫流调岗"
        assert jobs[0]["headcount"] == 2
        assert "预防医学" in jobs[0]["major_raw"]
        
        # 验证向上继承
        assert jobs[1]["unit_name"] == "杭州市疾病预防控制中心"
        assert jobs[1]["job_name"] == "理化检验岗"
        assert jobs[1]["headcount"] == 1
    finally:
        if os.path.exists(docx_path):
            os.remove(docx_path)

def test_pdf_parser_unified_interface():
    """测试 PdfJobParser parse_pdf 和 parse_file 解析表格及合计行过滤"""
    # 注册中文字体确保表格渲染
    font_path = "/usr/share/fonts/truetype/droid/DroidSansFallbackFull.ttf"
    font_name = "Helvetica"
    if os.path.exists(font_path):
        try:
            pdfmetrics.registerFont(TTFont("DroidSansFallback", font_path))
            font_name = "DroidSansFallback"
        except Exception:
            pass

    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tf:
        pdf_path = tf.name

    doc = SimpleDocTemplate(pdf_path, pagesize=letter)
    data = [
        ["用人单位", "岗位名称", "招考计划", "所需专业及代码", "学历要求", "备注"],
        ["浙江省疾控中心", "现场流调与应急", "1", "流行病与卫生统计学", "硕士", "具有公卫执业医师资格"],
        ["", "实验室理化检测", "2", "卫生检验与检疫", "本科", ""],
        ["总计", "总计", "3", "", "", ""]
    ]
    t = Table(data, style=[
        ("FONTNAME", (0, 0), (-1, -1), font_name),
        ("GRID", (0, 0), (-1, -1), 1, colors.black),
    ])
    doc.build([t])

    try:
        # 测试 parse_pdf 与 parse_file 两个统一接口
        jobs_pdf = PdfJobParser.parse_pdf(pdf_path)
        jobs_file = PdfJobParser.parse_file(pdf_path)
        assert len(jobs_pdf) == 2
        assert len(jobs_file) == 2
        assert jobs_pdf[0]["unit_name"] == "浙江省疾控中心"
        assert jobs_pdf[0]["job_name"] == "现场流调与应急"
        assert jobs_pdf[1]["unit_name"] == "浙江省疾控中心"
        assert jobs_pdf[1]["job_name"] == "实验室理化检测"
        
        # 测试通过 ParserDispatcher 分发
        disp_res = ParserDispatcher.parse_file(pdf_path)
        assert disp_res["status"] == "SUCCESS"
        assert len(disp_res["data"]) == 2
    finally:
        if os.path.exists(pdf_path):
            os.remove(pdf_path)

def test_html_table_parser():
    """测试 HTML 内嵌表格解析器"""
    html_content = """
    <div class="content">
      <p>招聘岗位如下：</p>
      <table border="1">
        <tr>
          <th>序号</th><th>用人单位</th><th>岗位名称</th><th>招聘计划</th><th>专业</th><th>学历</th>
        </tr>
        <tr>
          <td>1</td><td>宁波市疾控中心</td><td>职业卫生评价岗</td><td>1</td><td>劳动卫生与环境卫生学</td><td>硕士</td>
        </tr>
        <tr>
          <td colspan="6">合计：1人</td>
        </tr>
      </table>
    </div>
    """
    jobs = HtmlTableJobParser.parse_html_tables(html_content)
    assert len(jobs) == 1
    assert jobs[0]["unit_name"] == "宁波市疾控中心"
    assert jobs[0]["job_name"] == "职业卫生评价岗"
    assert jobs[0]["headcount"] == 1
    assert "劳动卫生" in jobs[0]["major_raw"]
