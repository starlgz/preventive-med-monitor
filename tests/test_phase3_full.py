import os, sys, asyncio
from app.parsers.column_mapper import ColumnMapper
from app.parsers.excel_parser import ExcelJobParser
from app.parsers.word_parser import WordJobParser
from app.parsers.html_table_parser import HtmlTableJobParser
from app.parsers.dispatcher import ParserDispatcher
import openpyxl, docx

def test_full_parsers():
    print("=== 1. 测试智能表头模糊映射 (ColumnMapper) ===")
    headers = ["序号", "招聘单位", "招考职位名称", "招聘计划人数", "所需专业与代码", "最低学历要求", "备注说明"]
    mapping = ColumnMapper.map_columns(headers)
    print(f"表头匹配结果: {mapping}")
    assert "unit_name" in mapping and "job_name" in mapping and "major" in mapping
    print("ColumnMapper 验证通过！\n")

    print("=== 2. 测试 Excel 岗位表解析 (含合并单元格向下自动填充) ===")
    os.makedirs("data/test_tmp", exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["序号", "招聘单位名称", "招聘岗位", "招聘人数", "学历要求", "专业要求"])
    ws.append([1, "杭州市疾病预防控制中心", "JK001", 2, "本科及以上", "预防医学、公共卫生"])
    ws.append([2, None, "JK002", 1, "硕士研究生", "流行病与卫生统计学"])
    ws.merge_cells("B2:B3") # 模拟合并单元格
    excel_path = "data/test_tmp/test_jobs.xlsx"
    wb.save(excel_path)

    excel_jobs = ExcelJobParser.parse_file(excel_path)
    print(f"Excel 提取到岗位数量: {len(excel_jobs)}")
    for j in excel_jobs:
        print(f"  - [{j['unit_name']}] 岗位:{j['job_name']}, 专业:{j['major_raw']}, 人数:{j['headcount']}")
    assert len(excel_jobs) == 2
    assert excel_jobs[1]["unit_name"] == "杭州市疾病预防控制中心" # 验证合并单元格自动填充
    print("Excel 解析器验证通过！\n")

    print("=== 3. 测试 Word 岗位表解析 (.docx) ===")
    doc = docx.Document()
    table = doc.add_table(rows=1, cols=4)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "用人单位"
    hdr_cells[1].text = "岗位名称"
    hdr_cells[2].text = "专业"
    hdr_cells[3].text = "学历"
    
    row_cells = table.add_row().cells
    row_cells[0].text = "宁波市疾控中心"
    row_cells[1].text = "理化检验岗"
    row_cells[2].text = "卫生检验与检疫、预防医学"
    row_cells[3].text = "本科"
    word_path = "data/test_tmp/test_jobs.docx"
    doc.save(word_path)

    word_jobs = WordJobParser.parse_file(word_path)
    print(f"Word 提取到岗位数量: {len(word_jobs)}")
    for j in word_jobs:
        print(f"  - [{j['unit_name']}] 岗位:{j['job_name']}, 专业:{j['major_raw']}")
    assert len(word_jobs) == 1
    print("Word 解析器验证通过！\n")

    print("=== 4. 测试 HTML 内嵌表格解析 ===")
    html_sample = """
    <div>
        <p>正文内容...</p>
        <table>
            <tr><th>单位</th><th>职位</th><th>专业要求</th><th>人数</th></tr>
            <tr><td>温州市疾病预防控制中心</td><td>公共卫生医师</td><td>预防医学、公共卫生</td><td>2</td></tr>
        </table>
    </div>
    """
    html_jobs = HtmlTableJobParser.parse_html_tables(html_sample)
    print(f"HTML 表格提取岗位数: {len(html_jobs)}")
    for j in html_jobs:
        print(f"  - [{j['unit_name']}] 岗位:{j['job_name']}, 专业:{j['major_raw']}")
    assert len(html_jobs) == 1
    print("HTML 表格解析器验证通过！\n")

    print("=== 5. 测试调度总中心 (ParserDispatcher) ===")
    async def test_dispatcher():
        res = await ParserDispatcher.parse_announcement(
            ann_id=999,
            title="测试综合招聘",
            raw_content=html_sample,
            raw_attachments_json="[]",
            unit_name="浙江疾控"
        )
        print(f"Dispatcher 返回岗位数: {res['jobs_count']}, 清洗后正文长度: {res['clean_text_length']}")
        assert res["jobs_count"] == 1
        print("ParserDispatcher 调度验证通过！\n")
    
    asyncio.run(test_dispatcher())

if __name__ == "__main__":
    test_full_parsers()
